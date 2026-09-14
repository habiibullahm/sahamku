"""Import official IDX disclosures from validated files or an optional data endpoint."""

from __future__ import annotations

import csv
import hashlib
import io
import json
import logging
import re
import sqlite3
from dataclasses import dataclass
from datetime import date, datetime
from pathlib import Path
from urllib.parse import urlparse
from urllib.request import Request, urlopen

from openpyxl import load_workbook

from sahamku import db
from sahamku.config import TZ, settings
from sahamku.universe import active_codes

log = logging.getLogger(__name__)
SOURCE = "idx_disclosures"
UA = "Mozilla/5.0 (compatible; SahamkuBot/1.0; official-disclosure-cache)"
CODE = re.compile(r"^[A-Z]{4}$")

ALIASES = {
    "ticker": {"ticker", "code", "kode", "kode_emiten", "kode_perusahaan"},
    "published": {
        "published", "published_at", "date", "tanggal", "tanggal_publikasi",
        "tanggal_pengumuman",
    },
    "category": {"category", "kategori", "jenis", "jenis_informasi"},
    "title": {"title", "judul", "perihal", "subject"},
    "url": {"url", "link", "tautan", "document_url", "download_url"},
    "document_id": {"document_id", "id", "nomor", "no_pengumuman"},
}


class DisclosureError(ValueError):
    """Invalid or unavailable disclosure input."""


@dataclass(frozen=True)
class RefreshResult:
    status: str
    imported: int = 0
    detail: str = ""


def _header(value: object) -> str:
    text = str(value or "").strip().lower()
    return re.sub(r"[^a-z0-9]+", "_", text).strip("_")


def _field_map(headers: list[object]) -> dict[str, int]:
    normalized = [_header(value) for value in headers]
    fields: dict[str, int] = {}
    for target, aliases in ALIASES.items():
        match = next((i for i, value in enumerate(normalized) if value in aliases), None)
        if match is not None:
            fields[target] = match
    missing = {"ticker", "published", "category", "title", "url"} - fields.keys()
    if missing:
        raise DisclosureError("kolom wajib tidak ada: " + ", ".join(sorted(missing)))
    return fields


def _rows_from_csv(data: bytes) -> list[dict[str, object]]:
    try:
        text = data.decode("utf-8-sig")
    except UnicodeDecodeError as exc:
        raise DisclosureError("CSV harus memakai encoding UTF-8") from exc
    try:
        dialect = csv.Sniffer().sniff(text[:4096], delimiters=",;\t|")
    except csv.Error:
        dialect = csv.excel
    rows = list(csv.reader(io.StringIO(text), dialect))
    return _mapped_rows(rows)


def _rows_from_xlsx(data: bytes) -> list[dict[str, object]]:
    try:
        workbook = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    except Exception as exc:
        raise DisclosureError("file XLSX tidak dapat dibaca") from exc
    sheet = workbook.active
    try:
        return _mapped_rows([list(row) for row in sheet.iter_rows(values_only=True)])
    finally:
        workbook.close()


def _mapped_rows(rows: list[list[object]]) -> list[dict[str, object]]:
    if not rows:
        raise DisclosureError("file disclosure kosong")
    fields = None
    header_index = 0
    for index, row in enumerate(rows[:20]):
        try:
            fields = _field_map(row)
            header_index = index
            break
        except DisclosureError:
            continue
    if fields is None:
        raise DisclosureError("header disclosure tidak ditemukan pada 20 baris pertama")
    output = []
    for row in rows[header_index + 1:]:
        if not any(value not in (None, "") for value in row):
            continue
        output.append({
            field: row[index] if index < len(row) else None
            for field, index in fields.items()
        })
    if not output:
        raise DisclosureError("file disclosure tidak memiliki data")
    return output


def _rows_from_json(data: bytes) -> list[dict[str, object]]:
    try:
        payload = json.loads(data)
    except (UnicodeDecodeError, json.JSONDecodeError) as exc:
        raise DisclosureError("respons JSON IDX tidak valid") from exc
    if isinstance(payload, dict):
        for key in ("data", "results", "items"):
            if isinstance(payload.get(key), list):
                payload = payload[key]
                break
    if not isinstance(payload, list) or not payload:
        raise DisclosureError("respons IDX tidak berisi daftar disclosure")
    if not all(isinstance(item, dict) for item in payload):
        raise DisclosureError("format item disclosure IDX tidak valid")
    headers = list(payload[0])
    fields = _field_map(headers)
    return [
        {field: item.get(headers[index]) for field, index in fields.items()}
        for item in payload
    ]


def load_bytes(data: bytes, filename: str = "disclosures.csv",
               content_type: str = "") -> list[dict[str, object]]:
    suffix = Path(filename).suffix.lower()
    if suffix == ".xlsx" or data.startswith(b"PK\x03\x04"):
        return _rows_from_xlsx(data)
    is_json = suffix == ".json" or "json" in content_type.lower()
    if is_json or data.lstrip().startswith((b"[", b"{")):
        return _rows_from_json(data)
    return _rows_from_csv(data)


def load_file(path: Path) -> list[dict[str, object]]:
    if path.suffix.lower() not in {".csv", ".xlsx", ".json"}:
        raise DisclosureError("format file harus CSV, XLSX, atau JSON")
    return load_bytes(path.read_bytes(), path.name)


def _official_url(value: object) -> str:
    url = str(value or "").strip()
    parsed = urlparse(url)
    host = (parsed.hostname or "").lower()
    if parsed.scheme != "https" or not (host == "idx.co.id" or host.endswith(".idx.co.id")):
        raise DisclosureError(f"URL disclosure bukan domain HTTPS resmi IDX: {url[:120]}")
    return url


def _published(value: object) -> str:
    if isinstance(value, datetime):
        parsed = value
    elif isinstance(value, date):
        parsed = datetime.combine(value, datetime.min.time())
    else:
        raw = str(value or "").strip()
        parsed = None
        try:
            parsed = datetime.fromisoformat(raw.replace("Z", "+00:00"))
        except ValueError:
            for pattern in ("%d/%m/%Y", "%d-%m-%Y", "%Y-%m-%d"):
                try:
                    parsed = datetime.strptime(raw, pattern)
                    break
                except ValueError:
                    continue
        if parsed is None:
            raise DisclosureError(f"tanggal disclosure tidak valid: {raw!r}")
    if parsed.tzinfo is None:
        parsed = parsed.replace(tzinfo=TZ)
    return parsed.astimezone(TZ).isoformat(timespec="minutes")


def normalize(records: list[dict[str, object]], conn: sqlite3.Connection,
              observed_at: datetime | None = None) -> list[dict[str, str]]:
    observed = (observed_at or datetime.now(TZ)).isoformat(timespec="minutes")
    master = conn.execute(
        "SELECT code FROM securities WHERE lower(status)='active' "
        "AND lower(instrument_type) IN ('stock','common stock','saham')"
    ).fetchall()
    known = {row["code"] for row in master} if master else set(active_codes(conn))
    output: list[dict[str, str]] = []
    seen: set[str] = set()
    errors: list[str] = []
    for line, row in enumerate(records, 2):
        code = str(row.get("ticker") or "").strip().upper().removesuffix(".JK")
        try:
            if not CODE.fullmatch(code):
                raise DisclosureError("kode emiten tidak valid")
            if known and code not in known:
                raise DisclosureError(f"{code} tidak ada di universe aktif")
            title = re.sub(r"\s+", " ", str(row.get("title") or "")).strip()
            category = re.sub(r"\s+", " ", str(row.get("category") or "")).strip()
            if not title or not category:
                raise DisclosureError("judul/kategori kosong")
            published = _published(row.get("published"))
            url = _official_url(row.get("url"))
            document_id = str(row.get("document_id") or "").strip()
            identity = (
                f"{code}|{document_id}" if document_id
                else f"{code}|{published}|{url}|{title}"
            )
            item_id = "idx-" + hashlib.sha256(identity.encode()).hexdigest()[:20]
            if item_id in seen:
                raise DisclosureError("disclosure duplikat")
            seen.add(item_id)
            output.append({
                "id": item_id, "ticker": code, "published": published,
                "category": category[:200], "title": title[:500], "url": url,
                "document_id": document_id[:100], "observed_at": observed,
            })
        except DisclosureError as exc:
            errors.append(f"baris {line}: {exc}")
    if errors:
        raise DisclosureError("; ".join(errors[:20]))
    return output


def ingest_records(conn: sqlite3.Connection, records: list[dict[str, object]],
                   observed_at: datetime | None = None) -> int:
    items = normalize(records, conn, observed_at)
    imported = 0
    for item in items:
        if db.idx_disclosure_upsert(conn, item):
            imported += 1
        db.news_insert(conn, {
            "id": item["id"], "source": "IDX", "title": item["title"],
            "summary": item["category"], "link": item["url"],
            "published": item["published"],
        }, item["ticker"])
    db.source_state_set(
        conn, SOURCE, "ok", f"{len(items)} valid; {imported} baru", success=True
    )
    return imported


def refresh(conn: sqlite3.Connection) -> RefreshResult:
    url = (settings.idx_disclosure_url or "").strip()
    if not url:
        return RefreshResult("disabled", detail="IDX_DISCLOSURE_URL belum diatur")
    conn.execute("SAVEPOINT idx_disclosure_refresh")
    try:
        _official_url(url)
        request = Request(
            url, headers={"User-Agent": UA, "Accept": "application/json,text/csv,*/*"}
        )
        with urlopen(request, timeout=settings.idx_disclosure_timeout_seconds) as response:
            _official_url(response.geturl())
            data = response.read(15_000_001)
            if len(data) > 15_000_000:
                raise DisclosureError("respons IDX melebihi 15 MB")
            disposition = response.headers.get("Content-Disposition", "")
            match = re.search(r'filename="?([^";]+)', disposition, re.IGNORECASE)
            filename = match.group(1) if match else Path(urlparse(url).path).name
            records = load_bytes(data, filename or "disclosures.json",
                                 response.headers.get("Content-Type", ""))
        imported = ingest_records(conn, records)
        conn.execute("RELEASE SAVEPOINT idx_disclosure_refresh")
        conn.commit()
        return RefreshResult("ok", imported, f"{len(records)} disclosure diterima")
    except Exception as exc:
        conn.execute("ROLLBACK TO SAVEPOINT idx_disclosure_refresh")
        conn.execute("RELEASE SAVEPOINT idx_disclosure_refresh")
        detail = f"{type(exc).__name__}: {exc}"[:500]
        db.source_state_set(conn, SOURCE, "stale", detail)
        conn.commit()
        log.warning("refresh disclosure IDX gagal; memakai cache: %s", detail)
        return RefreshResult("stale", detail=detail)
